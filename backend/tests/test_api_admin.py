import csv
import io
from datetime import UTC, datetime

import pytest
from openpyxl import load_workbook

from app.main import create_app

NOW = datetime(2026, 10, 23, 16, 0, tzinfo=UTC)  # Fri Oct 23, noon New York
FAKE_ID = "64b000000000000000000000"


def _admin_routes():
    app = create_app(start_background=False)
    for path, ops in app.openapi()["paths"].items():
        if path.startswith("/api/admin"):
            for method in ops:
                yield method.upper(), path.replace("{user_id}", FAKE_ID).replace("{session_id}", FAKE_ID)


ADMIN_ROUTES = list(_admin_routes())


def test_every_admin_route_is_listed():
    assert len(ADMIN_ROUTES) >= 14


@pytest.mark.parametrize("method,path", ADMIN_ROUTES)
async def test_general_user_gets_403_on_every_admin_route(user_client, method, path):
    r = await user_client.request(method, path, json={})
    assert r.status_code == 403, f"{method} {path} -> {r.status_code}"
    assert r.json()["error"]["code"] == "FORBIDDEN"


@pytest.mark.parametrize("method,path", ADMIN_ROUTES)
async def test_anonymous_gets_401_on_every_admin_route(client, method, path):
    r = await client.request(method, path, json={})
    assert r.status_code == 401, f"{method} {path} -> {r.status_code}"


async def _seed(make_user, make_client, clock):
    """Two volunteers in different Mandals with known sessions."""
    clock.set(NOW)
    await make_user("anjali.p", full_name="Anjali Patel", mandal_name="Edison", mandal_type="Yuvti")
    await make_user("meeras", full_name="Meera Shah", mandal_name="Jersey City", mandal_type="Kishori")
    async with make_client() as a:
        await a.login("anjali.p")
        # Oct 18 8:15 PM → Oct 19 12:45 AM (crosses midnight)
        await a.post("/api/me/sessions", json={"clock_in": "2026-10-18T20:15", "clock_out": "2026-10-19T00:45"})
        await a.post("/api/me/sessions", json={"clock_in": "2026-10-22T18:30", "clock_out": "2026-10-22T21:45"})
        await a.post("/api/me/sessions/clock-in")  # active since noon today
    async with make_client() as m:
        await m.login("meeras")
        await m.post("/api/me/sessions", json={"clock_in": "2026-10-22T16:00", "clock_out": "2026-10-22T20:30"})
        await m.post("/api/me/sessions", json={"clock_in": "2026-10-22T19:45", "clock_out": "2026-10-22T21:00", "confirm_overlap": True})
    clock.advance(hours=1)


async def test_stats(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    s = (await admin_client.get("/api/admin/stats")).json()
    assert s["total_volunteers"] == 2 and s["active_accounts"] == 2
    assert s["completed_sessions"] == 4
    assert s["total_hours"]["seconds"] == 16200 + 11700 + 16200 + 4500
    assert s["on_seva_now"] == 1
    assert s["hours_today"]["seconds"] == 3600  # live active session
    assert s["needs_review"]["total"] == 4 and s["needs_review"]["by_flag"]["overlap"] == 1
    edison = next(m for m in s["by_mandal"] if m["name"] == "Edison")
    assert edison["seconds"] == 16200 + 11700 + 3600


async def test_report_filters_and_clipped_totals(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    r = (await admin_client.get("/api/admin/sessions", params={"date": "2026-10-19"})).json()
    assert r["total"] == 1
    assert r["totals"]["duration"]["seconds"] == 2700  # only the after-midnight part counts on Oct 19
    assert r["items"][0]["duration"]["seconds"] == 16200  # row shows the full session

    r = (await admin_client.get("/api/admin/sessions", params={"mandal_name": "Jersey City"})).json()
    assert r["total"] == 2 and {i["volunteer"]["username"] for i in r["items"]} == {"meeras"}

    r = (await admin_client.get("/api/admin/sessions", params={"q": "ANJ"})).json()
    assert r["total"] == 3

    # "All" lists the active session but finalized totals exclude it.
    r = (await admin_client.get("/api/admin/sessions", params={"username": "anjali"})).json()
    assert r["total"] == 3 and r["totals"]["duration"]["seconds"] == 16200 + 11700
    r = (await admin_client.get("/api/admin/sessions", params={"status": "active"})).json()
    assert r["total"] == 1 and r["totals"]["duration"]["seconds"] == 3600

    r = (await admin_client.get("/api/admin/sessions", params={"status": "flagged", "mandal_type": "Kishori"})).json()
    assert r["total"] == 2


async def test_report_sorting_and_pagination(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    r = (await admin_client.get("/api/admin/sessions", params={"sort": "name", "order": "asc"})).json()
    assert [i["volunteer"]["full_name"] for i in r["items"]][:3] == ["Anjali Patel"] * 3
    r = (await admin_client.get("/api/admin/sessions", params={"sort": "date", "order": "desc"})).json()
    ins = [i["clock_in"] for i in r["items"]]
    assert ins == sorted(ins, reverse=True)
    assert (await admin_client.get("/api/admin/sessions", params={"page_size": 30})).status_code == 422
    r = (await admin_client.get("/api/admin/sessions", params={"page_size": 20, "page": 2})).json()
    assert r["items"] == [] and r["total"] == 5


async def test_totals_by_group(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    r = (await admin_client.get("/api/admin/totals", params={"group_by": "mandal_type", "preset": "month"})).json()
    by = {i["key"]: i["seconds"] for i in r["items"]}
    assert by == {"Kishori": 20700, "Yuvti": 27900, "Sahid": 0, "Sahradayi": 0}
    r = (await admin_client.get("/api/admin/totals", params={"group_by": "volunteer"})).json()
    assert r["items"][0]["label"] == "Anjali Patel"
    r = (await admin_client.get("/api/admin/totals", params={"group_by": "day", "from": "2026-10-18", "to": "2026-10-19"})).json()
    assert [(i["key"], i["seconds"]) for i in r["items"]] == [("2026-10-18", 13500), ("2026-10-19", 2700)]


async def test_export_csv_full_set_with_summary(admin_client, make_user, make_client, clock, db):
    await _seed(make_user, make_client, clock)
    await db.users.update_one({"username_lower": "meeras"}, {"$set": {"full_name": "=HYPERLINK(\"x\")"}})
    r = await admin_client.get("/api/admin/sessions/export", params={"format": "csv", "preset": "month"})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.text.lstrip("﻿"))))
    assert rows[0][:4] == ["Volunteer Name", "Username", "Mandal Name", "Mandal Type"]
    data = rows[1:6]
    assert len(data) == 5
    assert any(row[0] == "'=HYPERLINK(\"x\")" for row in data)  # formula injection neutralised
    flat = [",".join(r) for r in rows]
    assert any(line.startswith("Edison,") for line in flat)
    assert any(line.startswith("Kishori,") for line in flat)


async def test_export_xlsx(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    r = await admin_client.get("/api/admin/sessions/export", params={"format": "xlsx"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Sessions", "Summary"]
    assert wb["Sessions"].max_row == 6
    summary = [row for row in wb["Summary"].iter_rows(values_only=True)]
    assert any(row and row[0] == "Edison" for row in summary)


async def test_review_queue_correction_and_audit(admin_client, make_user, make_client, clock, db):
    await _seed(make_user, make_client, clock)
    q = (await admin_client.get("/api/admin/review", params={"type": "overlap"})).json()
    assert q["counts"]["overlap"] == 1 and len(q["items"]) == 1
    item = q["items"][0]
    assert item["volunteer"]["username"] == "meeras" and len(item["overlaps"]) == 1
    sid = item["id"]

    r = await admin_client.patch(f"/api/admin/sessions/{sid}", json={"clock_in": "2026-10-22T20:30"})
    assert r.status_code == 422  # reason required

    r = await admin_client.patch(
        f"/api/admin/sessions/{sid}", json={"clock_in": "2026-10-22T20:30", "reason": "Confirmed with Mandal coordinator"}
    )
    assert r.status_code == 200, r.text
    s = r.json()["session"]
    assert s["duration"]["seconds"] == 1800 and not s["needs_review"] and "overlap" not in s["flags"]

    log = (await admin_client.get("/api/admin/audit", params={"entity_id": sid})).json()["items"]
    assert log[0]["action"] == "correct" and log[0]["actor_role"] == "admin" and log[0]["actor_username"] == "organizer"
    assert log[0]["before"]["clock_in"] == "2026-10-22T23:45:00Z" and log[0]["after"]["clock_in"] == "2026-10-23T00:30:00Z"
    assert log[0]["reason"] == "Confirmed with Mandal coordinator" and log[0]["owner_name"] == "Meera Shah"


async def test_admin_closes_missing_clock_out(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    clock.advance(hours=12)
    q = (await admin_client.get("/api/admin/review", params={"type": "missing_clock_out"})).json()
    assert q["counts"]["missing_clock_out"] == 1
    sid = q["items"][0]["id"]
    r = await admin_client.patch(f"/api/admin/sessions/{sid}", json={"clock_out": "2026-10-23T15:30", "reason": "Confirmed"})
    assert r.status_code == 200, r.text
    s = r.json()["session"]
    assert s["status"] == "completed" and s["duration"]["seconds"] == int(3.5 * 3600)


async def test_mark_reviewed(admin_client, make_user, make_client, clock):
    await _seed(make_user, make_client, clock)
    q = (await admin_client.get("/api/admin/review", params={"type": "manual_entry"})).json()
    sid = q["items"][0]["id"]
    r = await admin_client.post(f"/api/admin/sessions/{sid}/mark-reviewed", json={})
    assert r.status_code == 200 and r.json()["session"]["needs_review"] is False
    log = (await admin_client.get("/api/admin/audit", params={"entity_id": sid})).json()["items"]
    assert log[0]["action"] == "review"


async def test_users_list_detail_and_deactivate(admin_client, make_user, make_client, clock, db):
    await _seed(make_user, make_client, clock)
    users = (await admin_client.get("/api/admin/users", params={"q": "meera"})).json()
    assert users["total"] == 1
    meera = users["items"][0]
    assert meera["lifetime"]["seconds"] == 16200 + 4500

    async with make_client() as m:
        await m.login("meeras")
        r = await admin_client.patch(f"/api/admin/users/{meera['id']}/status", json={"is_active": False})
        assert r.status_code == 200 and r.json()["user"]["is_active"] is False
        assert await db.refresh_tokens.count_documents({"user_id": __import__("bson").ObjectId(meera["id"]), "revoked_at": None}) == 0
        assert (await m.post("/api/auth/refresh")).status_code == 401
        assert (await m.get("/api/me/summary")).status_code == 401

    detail = (await admin_client.get(f"/api/admin/users/{meera['id']}")).json()
    assert detail["user"]["is_active"] is False and detail["lifetime"]["sessions"] == 2
    hist = (await admin_client.get(f"/api/admin/users/{meera['id']}/sessions")).json()
    assert hist["total"] == 2
    log = (await admin_client.get("/api/admin/audit", params={"entity_id": meera["id"]})).json()["items"]
    assert log[0]["action"] == "deactivate"

    filtered = (await admin_client.get("/api/admin/users", params={"is_active": "false"})).json()
    assert filtered["total"] == 1


async def test_admin_cannot_deactivate_self(admin_client, db):
    me = (await admin_client.get("/api/auth/me")).json()["user"]
    r = await admin_client.patch(f"/api/admin/users/{me['id']}/status", json={"is_active": False})
    assert r.status_code == 409
