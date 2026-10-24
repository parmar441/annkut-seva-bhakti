"""CSV and Excel export of the full filtered report, with totals by Mandal and Mandal Type."""

import csv
import io
import tempfile
from collections.abc import AsyncIterator, Iterator
from datetime import datetime

from motor.motor_asyncio import AsyncIOMotorDatabase
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from ..config import get_settings
from ..enums import MANDAL_NAMES, MANDAL_TYPES
from ..reports.timecalc import decimal_hours, ensure_utc, format_hm_short
from ..sessions.service import display_status

HEADERS = [
    "Volunteer Name",
    "Username",
    "Mandal Name",
    "Mandal Type",
    "Seva Date",
    "Clock-in Time",
    "Clock-out Time",
    "Session Duration",
    "Duration (hours)",
    "Duration (seconds)",
    "Session Status",
    "Flags",
]


def _safe(value):
    """Neutralise spreadsheet formula injection in text cells."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def session_row(doc: dict) -> list:
    tz = get_settings().tz
    cin = ensure_utc(doc["clock_in"]).astimezone(tz)
    cout = ensure_utc(doc["clock_out"]).astimezone(tz) if doc.get("clock_out") else None
    secs = doc.get("duration_seconds")
    v = doc["volunteer"]
    return [
        _safe(v["full_name"]),
        _safe(v["username"]),
        doc["mandal_name"],
        doc["mandal_type"],
        cin.strftime("%Y-%m-%d"),
        cin.strftime("%Y-%m-%d %I:%M %p"),
        cout.strftime("%Y-%m-%d %I:%M %p") if cout else "",
        format_hm_short(secs) if secs is not None else "",
        float(decimal_hours(secs)) if secs is not None else "",
        secs if secs is not None else "",
        display_status(doc).capitalize(),
        ", ".join(doc.get("flags", [])),
    ]


def summary_rows(label: str, total: dict, by_mandal: dict, by_type: dict, generated_at: datetime) -> list[list]:
    def line(name: str, t: dict) -> list:
        return [name, t["sessions"], format_hm_short(t["seconds"]), float(decimal_hours(t["seconds"]))]

    empty = {"seconds": 0, "sessions": 0}
    rows: list[list] = [
        ["Report", label],
        ["Generated", generated_at.astimezone(get_settings().tz).strftime("%Y-%m-%d %I:%M %p %Z")],
        ["Hours are clipped to the selected date range. Active sessions are excluded unless the Active filter is set."],
        [],
        ["Total", "Sessions", "Hours (h m)", "Hours"],
        line("All sessions", total),
        [],
        ["Mandal Name", "Sessions", "Hours (h m)", "Hours"],
        *[line(m, by_mandal.get(m, empty)) for m in MANDAL_NAMES],
        [],
        ["Mandal Type", "Sessions", "Hours (h m)", "Hours"],
        *[line(t, by_type.get(t, empty)) for t in MANDAL_TYPES],
    ]
    return rows


def _csv_line(row: list) -> str:
    buf = io.StringIO()
    csv.writer(buf).writerow(row)
    return buf.getvalue()


async def stream_csv(rows: AsyncIterator[dict], summary: list[list]) -> AsyncIterator[str]:
    yield "﻿"  # BOM so Excel opens UTF-8 correctly
    yield _csv_line(HEADERS)
    async for doc in rows:
        yield _csv_line(session_row(doc))
    yield "\r\n"
    for r in summary:
        yield _csv_line(r)


async def build_xlsx(rows: AsyncIterator[dict], summary: list[list]) -> tempfile.SpooledTemporaryFile:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sessions")
    bold = Font(bold=True)
    def header(cells: list) -> list:
        out = []
        for c in cells:
            cell = WriteOnlyCell(ws, value=c)
            cell.font = bold
            out.append(cell)
        return out

    ws.append(header(HEADERS))
    async for doc in rows:
        ws.append(session_row(doc))

    ss = wb.create_sheet("Summary")
    for r in summary:
        ss.append(r)

    f = tempfile.SpooledTemporaryFile(max_size=10 * 1024 * 1024)
    wb.save(f)
    f.seek(0)
    return f


def iter_file(f, chunk: int = 64 * 1024) -> Iterator[bytes]:
    try:
        while data := f.read(chunk):
            yield data
    finally:
        f.close()


async def iter_report_rows(db: AsyncIOMotorDatabase, pipeline: list[dict]) -> AsyncIterator[dict]:
    async for doc in db.seva_sessions.aggregate(pipeline):
        yield doc
